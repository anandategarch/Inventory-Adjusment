"""
accurate_bot.py — logika otomasi Accurate Online (v4.7, pasangan ui_app.py v4.7).
Flow terbukti: isi tanggal -> info -> upload -> konfirmasi -> simpan.
Verifikasi ketat: tanggal, akun (harus sama target), cabang (reset + sama target).
Pop-up: panduan upload diabaikan; hasil/error/stok diproses; LANJUTKAN otomatis.
v4.7: database COA & Keterangan dibaca otomatis dari folder aplikasi.
"""
import os
import re
import socket
import subprocess
import time
import random

from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC

BOT_VERSION = "4.7"
DEBUG_PORT = 9222

MODE_APPROVE = "APPROVE"
MODE_DRAFT = "DRAFT"
MODE_IMPORT = "HANYA_IMPORT"

FILE_SKIP_KEYWORDS = ("RAW MATERIAL", "DEVIASI", "ADJUSTMENT STOCK")

FILE_KEYWORD_TO_MEMO_PREFIX = {
    "BEBAN ATK": "BEBAN ATK",
    "BEBAN GAS": "BEBAN GAS",
    "BEBAN KEBERSIHAN": "BEBAN KEBERSIHAN",
    "CONSUME": "COM (CONSUME)",
    "PACKAGING USAGE": "PACKAGING",
    "BIAYA BAHAN PEMBANTU": "BIAYA BAHAN PEMBANTU",
    "BIAYA PERLENGKAPAN": "BIAYA PERLENGKAPAN",
    "BIAYA SERAGAM": "BIAYA SERAGAM",
    "SPAREPART INVENTARIS": "SPAREPART INVENTARIS",
}

ADJUSTMENT_MARKERS = ("ADJUSTMENT STOCK", "(DEVIASI)", "DEVIASI")


def norm(s):
    return re.sub(r"\s+", " ", (s or "").strip().upper())


def fmt_coa(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _to_text(v):
    return "" if v is None else str(v).strip()


def split_suffix(text):
    parts = (text or "").split("_")
    if len(parts) >= 3:
        return "_".join(parts[:-2]), parts[-2], parts[-1]
    return text, None, None


def human_pause(min_s=0.4, max_s=1.2):
    time.sleep(random.uniform(min_s, max_s))


# ============================================================
# CHROME
# ============================================================
def is_port_open(port):
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        s.settimeout(0.5)
        return s.connect_ex(("127.0.0.1", port)) == 0


def launch_chrome_silent():
    candidates = [
        os.path.expandvars(r"%ProgramFiles%\Google\Chrome\Application\chrome.exe"),
        os.path.expandvars(r"%ProgramFiles(x86)%\Google\Chrome\Application\chrome.exe"),
        os.path.expandvars(r"%LocalAppData%\Google\Chrome\Application\chrome.exe"),
    ]
    chrome = next((p for p in candidates if os.path.exists(p)), None)
    if not chrome:
        return False
    user_data = os.path.expandvars(r"%TEMP%\ChromeDebug_Accurate")
    try:
        subprocess.Popen(
            [chrome, f"--remote-debugging-port={DEBUG_PORT}",
             f"--user-data-dir={user_data}", "https://accurate.id"],
            creationflags=0x08000000
        )
        return True
    except Exception:
        return False


# ============================================================
# DATABASE (per file Excel)
# ============================================================
def _enrich(coa, memo):
    core, cab, period = split_suffix(memo)
    core_key = norm(core[3:]) if core.upper().startswith("IA ") else norm(core)
    return {"coa": coa, "memo": memo, "nm": norm(memo),
            "core": core, "core_nm": norm(core), "core_key": core_key,
            "cab": cab, "period": period}


def load_database_map(db_path):
    try:
        import openpyxl
    except ImportError:
        openpyxl = None
    try:
        import xlrd
    except ImportError:
        xlrd = None

    ext = os.path.splitext(db_path)[1].lower()
    result = {}

    if ext == ".xlsx":
        if openpyxl is None:
            raise RuntimeError("Library 'openpyxl' belum terpasang. Jalankan: python -m pip install openpyxl")
        wb = openpyxl.load_workbook(db_path, data_only=True)
        try:
            for ws in wb.worksheets:
                rows = []
                cabang_db = _to_text(ws["C1"].value)
                for r in ws.iter_rows(min_row=2, values_only=True):
                    coa = fmt_coa(r[1] if len(r) > 1 else None)
                    memo = _to_text(r[2] if len(r) > 2 else None)
                    if coa and memo:
                        rows.append(_enrich(coa, memo))
                if rows:
                    result[ws.title] = {"cabang_db": cabang_db, "rows": rows}
        finally:
            wb.close()
    elif ext == ".xls":
        if xlrd is None:
            raise RuntimeError("Library 'xlrd' belum terpasang. Jalankan: python -m pip install xlrd")
        wb = xlrd.open_workbook(db_path)
        for idx in range(wb.nsheets):
            ws = wb.sheet_by_index(idx)
            rows = []
            cabang_db = _to_text(ws.cell_value(0, 2))
            for i in range(1, ws.nrows):
                coa = fmt_coa(ws.cell_value(i, 1))
                memo = _to_text(ws.cell_value(i, 2))
                if coa and memo:
                    rows.append(_enrich(coa, memo))
            if rows:
                result[ws.name] = {"cabang_db": cabang_db, "rows": rows}
    else:
        raise ValueError(f"Format tidak didukung: {ext}")

    if not result:
        raise ValueError("Database tidak memiliki baris COA + Keterangan yang terisi di sheet mana pun.")
    return result


# ============================================================
# DATABASE DARI FOLDER APLIKASI (fitur v4.7)
# ============================================================
DB_FOLDER_NAME = "Database COA&Keterangan"


def get_db_folder():
    """Folder database berada di dalam folder aplikasi."""
    base = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base, DB_FOLDER_NAME)


def load_database_folder(folder_path):
    """Scan folder database; gabungkan semua sheet berisi COA dari semua file Excel.
    Return (db_map, meta)."""
    if not os.path.isdir(folder_path):
        try:
            os.makedirs(folder_path, exist_ok=True)
            return {}, {"files": [], "total_rows": 0, "created": True}
        except Exception:
            return {}, {"files": [], "total_rows": 0, "created": False}

    db_map = {}
    files_used = []
    total_rows = 0
    for fn in sorted(os.listdir(folder_path)):
        if fn.startswith("~$"):
            continue
        if not fn.lower().endswith((".xlsx", ".xls")):
            continue
        path = os.path.join(folder_path, fn)
        try:
            sub = load_database_map(path)
        except Exception:
            continue
        for sheet, d in sub.items():
            key = f"{os.path.splitext(fn)[0]} | {sheet}"
            db_map[key] = d
            files_used.append(fn)
            total_rows += len(d["rows"])
    return db_map, {"files": sorted(set(files_used)), "total_rows": total_rows, "created": False}


# ============================================================
# PARSER + MATCHING
# ============================================================
def parse_file_info(filename):
    base = os.path.splitext(filename)[0]
    m = re.match(r"^(.*?)_\d+\.", base)
    keyword = (m.group(1) if m else base.split("_")[0]).strip()
    branch = None
    parts = base.split("_")
    if len(parts) >= 2:
        seg = parts[-2]
        if "." in seg:
            seg = seg.split(".")[-1]
        branch = seg.strip() or None
    return keyword, branch, base


def _rebuild(row, branch):
    if branch and row["period"]:
        return f"{row['core']}_{branch}_{row['period']}"
    return row["memo"]


def _is_adjustment(row):
    return any(mk in row["memo"].upper() for mk in ADJUSTMENT_MARKERS)


def match_file(db_map, base, keyword, branch):
    nb = norm(base)
    ncf = norm(split_suffix(base)[0])

    sheets = list(db_map.keys())
    prefer = [s for s in sheets if norm(db_map[s]["cabang_db"]) == norm(branch or "")]
    order = prefer + [s for s in sheets if s not in prefer]

    for name in order:
        for r in db_map[name]["rows"]:
            if r["nm"] == nb:
                return {"coa": r["coa"], "memo": r["memo"]}, name, "EXACT-FULL"

    for name in order:
        for r in db_map[name]["rows"]:
            if _is_adjustment(r):
                continue
            if r["core_nm"] == ncf:
                return {"coa": r["coa"], "memo": _rebuild(r, branch)}, name, "EXACT-CORE"

    prefix = FILE_KEYWORD_TO_MEMO_PREFIX.get(norm(keyword))
    if prefix:
        for name in order:
            for r in db_map[name]["rows"]:
                if _is_adjustment(r):
                    continue
                if r["core_key"].startswith(prefix):
                    return {"coa": r["coa"], "memo": _rebuild(r, branch)}, name, "KEYWORD"
        return None, None, "baris cocok tidak ditemukan (mungkin baris adjustment/deviasi)"
    return None, None, f"kata kunci '{keyword}' tidak dikenal"


def build_file_plan(db_map, files):
    plan = []
    for path in files:
        fn = os.path.basename(path)
        keyword, branch, base = parse_file_info(fn)
        folder_rel = os.path.basename(os.path.dirname(path))
        base_item = {"path": path, "folder": folder_rel, "filename": fn,
                     "keyword": keyword, "branch": branch}

        if any(s in norm(base) for s in FILE_SKIP_KEYWORDS):
            plan.append({**base_item, "coa": "", "memo": "", "status": "DILEWATI",
                         "reason": "masuk daftar skip (Raw Material/Deviasi/Adjustment)"})
            continue
        if not db_map:
            plan.append({**base_item, "coa": "", "memo": "", "status": "MENUNGGU DB",
                         "reason": "database belum dipilih"})
            continue

        res, sheet, how = match_file(db_map, base, keyword, branch)
        if res:
            plan.append({**base_item, "coa": res["coa"], "memo": res["memo"],
                         "status": "OK", "reason": f"{how} @ {sheet}"})
        else:
            plan.append({**base_item, "coa": "", "memo": "", "status": "TIDAK COCOK",
                         "reason": how})
    return plan


def collect_excel_files(root_folder):
    folders_with_files, total_files = [], 0
    for dirpath, dirnames, filenames in os.walk(root_folder):
        dirnames.sort()
        excel_files = sorted([os.path.join(dirpath, f) for f in filenames
                              if (f.lower().endswith(".xlsx") or f.lower().endswith(".xls")) and not f.startswith("~$")])
        if excel_files:
            rel = os.path.relpath(dirpath, root_folder)
            folders_with_files.append({"path": dirpath, "name": "(Folder Utama)" if rel == "." else rel, "files": excel_files})
            total_files += len(excel_files)
    return folders_with_files, total_files


# ============================================================
# SELENIUM HELPERS
# ============================================================
def find_and_switch_iframe(driver, selector, timeout=15):
    start = time.time()
    while time.time() - start < timeout:
        driver.switch_to.default_content()
        if len(driver.find_elements(By.CSS_SELECTOR, selector)) > 0:
            return True
        for iframe in driver.find_elements(By.TAG_NAME, "iframe"):
            try:
                driver.switch_to.default_content()
                driver.switch_to.frame(iframe)
                if len(driver.find_elements(By.CSS_SELECTOR, selector)) > 0:
                    return True
            except Exception:
                continue
        driver.switch_to.default_content()
        time.sleep(0.5)
    driver.switch_to.default_content()
    return False


def set_transaction_date(driver, date_str):
    return driver.execute_script("""
        let inp = document.querySelector('input[name="transDate"]') ||
                  document.querySelector('input[name*="Date"]') ||
                  document.querySelector('input[name*="date"]');
        if (!inp) return null;
        inp.focus();
        let nativeSet = Object.getOwnPropertyDescriptor(HTMLInputElement.prototype, 'value')?.set;
        if (nativeSet) nativeSet.call(inp, arguments[0]); else inp.value = arguments[0];
        ['input','change','blur'].forEach(ev => inp.dispatchEvent(new Event(ev, {bubbles:true})));
        return inp.value || null;
    """, date_str)


def smart_detect_accurate_tab(driver, total_timeout=25):
    FORM_SELECTORS = ('button[name="btnGetFrom"], input[name="adjustmentAccount"], '
                      'input[name="branch"], input[name="transDate"], '
                      'a[title="Rincian Barang"], a[title="Info lainnya"]')

    def get_accurate_tabs():
        result = []
        for h in driver.window_handles:
            try:
                driver.switch_to.window(h)
                url = (driver.current_url or '').lower()
                title = (driver.title or '').lower()
                if 'accurate' in url or 'accurate' in title:
                    score = 0
                    if 'item-adjustment' in url: score += 100
                    elif 'adjustment' in url: score += 80
                    elif 'inventory' in url: score += 60
                    else: score += 20
                    if 'penyesuaian' in title: score += 40
                    result.append({'handle': h, 'url': driver.current_url, 'title': driver.title, 'score': score})
            except Exception:
                continue
        result.sort(key=lambda x: x['score'], reverse=True)
        return result

    def form_exists():
        try:
            driver.switch_to.default_content()
            if len(driver.find_elements(By.CSS_SELECTOR, FORM_SELECTORS)) > 0:
                return True
            for iframe in driver.find_elements(By.TAG_NAME, 'iframe'):
                try:
                    driver.switch_to.default_content()
                    driver.switch_to.frame(iframe)
                    if len(driver.find_elements(By.CSS_SELECTOR, FORM_SELECTORS)) > 0:
                        driver.switch_to.default_content()
                        return True
                except Exception:
                    continue
            driver.switch_to.default_content()
        except Exception:
            pass
        return False

    deadline = time.time() + total_timeout
    best_tab = None
    while time.time() < deadline:
        tabs = get_accurate_tabs()
        if not tabs:
            time.sleep(1.5)
            continue
        best_tab = tabs[0]
        for tab in tabs:
            try:
                driver.switch_to.window(tab['handle'])
                if form_exists():
                    driver.switch_to.default_content()
                    return {'found': True, 'reason': 'ok', 'tab': tab}
            except Exception:
                continue
        time.sleep(2)

    if best_tab is None:
        return {'found': False, 'tab': None,
                'message': "Tidak ada tab Accurate yang terbuka di Chrome.\n\nBuka https://accurate.id, login, lalu coba lagi."}
    return {'found': False, 'tab': best_tab,
            'message': (f"Tab Accurate ditemukan: \"{best_tab['title'][:50]}\"\n"
                        "TAPI form Penyesuaian Persediaan belum terbuka.\n\n"
                        "Solusi: Klik 'Buat Baru' di Penyesuaian Persediaan,\n"
                        "tunggu form terbuka, lalu klik tombol mode lagi.")}


# ============================================================
# FORM FILLING (verifikasi ketat)
# ============================================================
def check_account_observable(driver):
    return driver.execute_script("""
        let accInp = document.querySelector('input[name="adjustmentAccount"]');
        if (!accInp) return {filled: false, msg: 'input tidak ditemukan'};
        try {
            let ctx = ko.contextFor(accInp);
            let vm  = ctx.$data;
            if (vm && vm.formData && typeof vm.formData.adjustmentAccount === 'function') {
                let val = vm.formData.adjustmentAccount();
                if (val && val !== null) return { filled: true, no: val.no || '', name: val.name || '' };
            }
        } catch(e) {}
        return {filled: false, msg: 'masih kosong'};
    """)


def wait_account_filled(driver, timeout=6.0):
    end = time.time() + timeout
    while time.time() < end:
        res = check_account_observable(driver)
        if res and res.get("filled"):
            return res
        time.sleep(0.4)
    return check_account_observable(driver)


def verify_form_observables(driver):
    return driver.execute_script("""
        function getVM(el) {
            try { return ko.contextFor(el).$data; } catch(e) {}
            try { return ko.dataFor(el); } catch(e) {}
            return null;
        }
        let out = {accountFound: false, branchFound: false, accountNo: null, branchName: null};
        let accInp = document.querySelector('input[name="adjustmentAccount"]');
        if (accInp) {
            let vm = getVM(accInp);
            if (vm && vm.formData && typeof vm.formData.adjustmentAccount === 'function') {
                let v = vm.formData.adjustmentAccount();
                if (v) { out.accountFound = true; out.accountNo = v.no || v.name || 'terisi'; }
            }
        }
        let brInp = document.querySelector('input[name="branch"]');
        if (brInp) {
            let vm = getVM(brInp);
            if (vm && vm.formData) {
                if (typeof vm.formData.branch === 'function') {
                    let v = vm.formData.branch();
                    if (v) { out.branchFound = true; out.branchName = v.name || ''; }
                } else if (typeof vm.formData.branchId === 'function') {
                    if (vm.formData.branchId()) out.branchFound = true;
                }
            }
        }
        return out;
    """)


def fill_info_lainnya(driver, account_search, memo_text, branch_text, log):
    driver.execute_script("""
        let tabInfo = document.querySelector('a[title="Info lainnya"], a.left-tab[title="Info lainnya"], .icn-transaction-header');
        if (tabInfo) {
            let target = tabInfo.closest('a') || tabInfo;
            ['mousedown','mouseup','click'].forEach(e => target.dispatchEvent(new MouseEvent(e, {bubbles:true, cancelable:true})));
        }
    """)
    human_pause(0.8, 1.6)

    driver.execute_script("""
        let textareas = document.querySelectorAll('textarea');
        textareas.forEach(ta => {
            let nativeSet = Object.getOwnPropertyDescriptor(HTMLTextAreaElement.prototype, 'value')?.set;
            if (nativeSet) nativeSet.call(ta, arguments[0]); else ta.value = arguments[0];
            ['input','change','blur'].forEach(ev => ta.dispatchEvent(new Event(ev, {bubbles:true})));
        });
    """, memo_text)
    human_pause(0.4, 0.9)

    chosen = driver.execute_script("""
        let inp = document.querySelector('input[name="branch"]');
        if (!inp || !window.ko) return null;
        let vm = ko.dataFor(inp);
        if (!vm || !vm.formData) return null;
        if (typeof vm.formData.branchId === 'function') vm.formData.branchId(null);
        if (typeof vm.formData.branch === 'function') vm.formData.branch(null);
        if (!window.acc || !acc.staticData || typeof acc.staticData.branchListOption !== 'function') return null;
        let optionsList = acc.staticData.branchListOption();
        let target = arguments[0].toUpperCase();
        let matchObj = optionsList.find(o => o && o.name && o.name.toUpperCase() === target);
        if (!matchObj) matchObj = optionsList.find(o => o && o.name && o.name.toUpperCase().includes(target));
        if (!matchObj) return null;
        if (typeof vm.formData.branchId === 'function') vm.formData.branchId(matchObj.id);
        if (typeof vm.formData.branch === 'function') vm.formData.branch(matchObj);
        return matchObj.name || null;
    """, branch_text)
    if chosen:
        log(f"Cabang diset ke: {chosen}")
    else:
        log(f"Cabang '{branch_text}' tidak ditemukan di daftar Accurate.", "ERROR")
    human_pause(0.4, 0.9)

    driver.execute_script("""
        let accInp = document.querySelector('input[name="adjustmentAccount"]');
        if (!accInp) return;
        try {
            let ctx = ko.contextFor(accInp); let vm = ctx.$data;
            if (vm && vm.formData && typeof vm.formData.adjustmentAccount === 'function') vm.formData.adjustmentAccount(null);
            if (vm && typeof vm.searchKeywordAccount === 'function') vm.searchKeywordAccount('');
        } catch(e) {}
        let container = accInp.closest('.lookupbox') || accInp.closest('.input-control') || accInp.parentElement;
        if (container) {
            let btns = container.querySelectorAll('button, .icon-cancel-2, .icn-close');
            btns.forEach(b => {
                if (b.classList.contains('icon-cancel-2') || b.classList.contains('icn-close') ||
                    (b.getAttribute('data-bind') && b.getAttribute('data-bind').includes('removeItem'))) {
                    let target = (b.tagName === 'I' ? (b.parentElement || b) : b);
                    if (target.isConnected) ['mousedown','mouseup','click'].forEach(e => target.dispatchEvent(new MouseEvent(e, {bubbles:true, cancelable:true})));
                }
            });
        }
    """)
    human_pause(0.5, 1.0)

    driver.execute_script("""
        let accInp = document.querySelector('input[name="adjustmentAccount"]');
        if (!accInp) return;
        accInp.style.display = 'inline-block'; accInp.style.visibility = 'visible'; accInp.focus();
        let nativeSet = Object.getOwnPropertyDescriptor(HTMLInputElement.prototype, 'value')?.set;
        if (nativeSet) nativeSet.call(accInp, arguments[0]); else accInp.value = arguments[0];
        ['input','change','keydown','keypress','keyup'].forEach(ev => accInp.dispatchEvent(new Event(ev, {bubbles:true})));
        try {
            let ctx = ko.contextFor(accInp); let vm = ctx.$data;
            if (vm && typeof vm.searchKeywordAccount === 'function') vm.searchKeywordAccount(arguments[0]);
        } catch(e) {}
        let container = accInp.closest('.lookupbox') || accInp.closest('.input-control') || accInp.parentElement;
        if (container) {
            let searchBtn = container.querySelector('.btn-search, button, .icon-search');
            if (searchBtn && searchBtn.isConnected) ['mousedown','mouseup','click'].forEach(e => searchBtn.dispatchEvent(new MouseEvent(e, {bubbles:true, cancelable:true})));
        }
        ['keydown','keypress','keyup'].forEach(type => accInp.dispatchEvent(new KeyboardEvent(type, {key:'Enter', keyCode:13, which:13, bubbles:true})));
    """, account_search)

    acc_res = wait_account_filled(driver, timeout=3.0)
    if not acc_res.get("filled"):
        driver.execute_script("""
            let selectors = ['.lookup-popup li', '.lookup-popup tr', '.lookup-popup a', '.autocomplete-suggestions div',
                '.dropdown-menu li', '.dropdown-menu a', 'ul.dropdown-menu li', '.suggestion-item', '.lookup-result li',
                '.lookup-result tr', '[class*="lookup"] li', '[class*="lookup"] tr', '[class*="dropdown"] li',
                '[class*="popup"] li', '[class*="popup"] tr', '[class*="suggest"] li', '[class*="suggest"] div',
                '[class*="list-item"]', '.ui-lookup-result li', '.ui-lookup-result tr', '[class*="search-result"] li',
                '[class*="search-result"] tr'].join(', ');
            let items = Array.from(document.querySelectorAll(selectors)).filter(el => {
                let r = el.getBoundingClientRect(); return r.width > 0 && r.height > 0; });
            let target = arguments[0].toUpperCase();
            let match = items.find(i => (i.innerText || i.textContent || '').toUpperCase().includes(target));
            if (match && match.isConnected) ['mousedown','mouseup','click'].forEach(e => match.dispatchEvent(new MouseEvent(e, {bubbles:true, cancelable:true})));
            else if (items.length > 0 && items[0].isConnected) ['mousedown','mouseup','click'].forEach(e => items[0].dispatchEvent(new MouseEvent(e, {bubbles:true, cancelable:true})));
        """, account_search)
        acc_res = wait_account_filled(driver, timeout=5.0)

    account_match = bool(acc_res.get("filled")) and norm(acc_res.get("no", "")) == norm(account_search)
    if acc_res.get("filled") and not account_match:
        log(f"AKUN TIDAK SESUAI TARGET: terpilih {acc_res.get('no')} vs target {account_search}. File digagalkan.", "ERROR")
    elif acc_res.get("filled"):
        log(f"Akun terisi: {acc_res.get('no')}", "SUCCESS")
    else:
        log("Akun belum terisi menurut observable.", "WARN")

    human_pause(0.5, 1.0)
    verify = verify_form_observables(driver)
    verify["accountMatch"] = account_match
    verify["branchSet"] = bool(chosen) and norm(chosen) == norm(branch_text)
    return verify


# ============================================================
# IMPORT
# ============================================================
def trigger_excel_import(driver, file_path):
    driver.execute_script("""
        function triggerClick(el) {
            if (!el || !el.isConnected) return;
            ['pointerdown','mousedown','pointerup','mouseup','click'].forEach(evt => {
                try { el.dispatchEvent(new MouseEvent(evt, {bubbles:true, cancelable:true, view:window, buttons:1})); } catch(e) {}
            });
            if (typeof el.click === 'function') { try { el.click(); } catch(e) {} }
        }
        let tabRincian = document.querySelector('a[title="Rincian Barang"], a.left-tab[title="Rincian Barang"], .icn-transaction-detail');
        if (tabRincian) triggerClick(tabRincian.closest('a') || tabRincian);
    """)
    human_pause(1.0, 2.0)

    find_and_switch_iframe(driver, 'button[name="btnGetFrom"]', timeout=5)
    btn = WebDriverWait(driver, 10).until(EC.presence_of_element_located((By.CSS_SELECTOR, 'button[name="btnGetFrom"]')))
    driver.execute_script("arguments[0].click();", btn)
    human_pause(0.8, 1.6)

    link = WebDriverWait(driver, 10).until(EC.presence_of_element_located((By.CSS_SELECTOR, 'a[data-bind*="getFromExcel"]')))
    driver.execute_script("arguments[0].click();", link)
    human_pause(1.2, 2.4)

    file_inp = WebDriverWait(driver, 10).until(EC.presence_of_element_located((By.CSS_SELECTOR, 'input[type="file"]')))
    file_inp.send_keys(file_path)
    human_pause(2.0, 4.0)


# ============================================================
# POP-UP
# ============================================================
def handle_popups(driver, strict=False):
    return driver.execute_script("""
        function triggerClick(el) {
            if (!el || !el.isConnected) return;
            ['pointerdown','mousedown','pointerup','mouseup','click'].forEach(evt => {
                try { el.dispatchEvent(new MouseEvent(evt, {bubbles:true, cancelable:true, view:window, buttons:1})); } catch(e) {}
            });
            if (typeof el.click === 'function') { try { el.click(); } catch(e) {} }
        }
        function evalDoc(doc, strictMode) {
            let out = {found:false, clicked:[], text:"", isSuccess:false, isError:false, isStockWarning:false};
            let dialogs = doc.querySelectorAll('.grid.fluid.response, .window, .dialog, .message-box, .notify, div[class*="popup"], div[id*="dialog"]');
            dialogs.forEach(dlg => {
                let r = dlg.getBoundingClientRect();
                if (!r || r.width <= 0 || r.height <= 0) return;
                let t = (dlg.innerText || dlg.textContent || '').trim();
                if (!t) return;
                let up = t.toUpperCase();
                let isResult  = up.includes('BERHASIL TERIMPOR') || up.includes('GAGAL IMPOR');
                let isErrorP  = up.includes('PERMASALAHAN') || up.includes('PERBAIKI');
                let isStock   = up.includes('TIDAK MENCUKUPI') || up.includes('STOK BARANG');
                let isGuide   = (up.includes('TEMPLATE FILE EXCEL') || up.includes('PILIH FILE EXCEL') || up.includes('UNGGAH FILE EXCEL'))
                                && !isResult && !isErrorP && !isStock;
                if (isGuide) return;
                let isKnown = isResult || isErrorP || isStock;
                if (strictMode && !isKnown) return;
                out.found = true;
                out.text += t + " ";
                if (isStock) out.isStockWarning = true;
                if (isResult) out.isSuccess = true;
                if (!isResult && isErrorP) out.isError = true;

                let btnContinue = dlg.querySelector('button[name="btnContinue"]');
                if (btnContinue) { triggerClick(btnContinue); out.clicked.push('LANJUTKAN'); return; }
                let btns = Array.from(dlg.querySelectorAll('button, a, input[type="button"]'));
                let label = b => (b.innerText || b.value || b.textContent || '').trim().toUpperCase();
                let okBtn = btns.find(b => {
                    let lb = label(b);
                    return lb === 'OK' || lb === 'TUTUP' || lb === 'CLOSE' || lb === 'X' || lb === '×' || b.classList.contains('close');
                });
                if (okBtn) { triggerClick(okBtn); out.clicked.push('OK'); }
            });
            return out;
        }
        let out = {found:false, clicked:[], text:"", isSuccess:false, isError:false, isStockWarning:false};
        function merge(r) {
            if (!r || !r.found) return;
            out.found = true;
            out.text += r.text + " ";
            out.isStockWarning = out.isStockWarning || r.isStockWarning;
            out.isSuccess = out.isSuccess || r.isSuccess;
            out.isError   = out.isError   || r.isError;
            out.clicked   = out.clicked.concat(r.clicked);
        }
        merge(evalDoc(document, arguments[0]));
        if (window.top !== window) { try { merge(evalDoc(window.top.document, arguments[0])); } catch(e) {} }
        let iframes = document.querySelectorAll('iframe');
        for (let f of iframes) {
            try { if (f.contentDocument) merge(evalDoc(f.contentDocument, arguments[0])); } catch(e) {}
        }
        return out;
    """, bool(strict))


def wait_import_result(driver, timeout=45, log=None):
    deadline = time.time() + timeout
    while time.time() < deadline:
        human_pause(0.8, 1.6)
        res = handle_popups(driver, strict=True)
        if res and res.get("found"):
            if log and res.get("clicked"):
                kind = "WARN" if res.get("isStockWarning") else "INFO"
                log(f"Pop-up hasil impor terdeteksi -> klik: {', '.join(res['clicked'])}", kind)
            return res
    return None


def handle_popups_flow(driver, wait_rounds=20, drain_rounds=8, delay=0.6, log=None):
    last = None
    first = None
    for _ in range(wait_rounds):
        time.sleep(random.uniform(delay, delay + 0.7))
        res = handle_popups(driver, strict=False)
        last = res
        if res and res.get("found"):
            first = res
            break
    if first is None:
        return last, []

    if log and first.get("clicked"):
        kind = "WARN" if first.get("isStockWarning") else "INFO"
        log(f"Pop-up terdeteksi -> klik: {', '.join(first['clicked'])}", kind)

    merged = dict(first)
    clicked = list(first.get("clicked", []))
    cur = first
    for _ in range(drain_rounds):
        if not cur or not cur.get("clicked"):
            break
        time.sleep(random.uniform(delay, delay + 0.7))
        res = handle_popups(driver, strict=False)
        if not res or not res.get("found"):
            break
        clicked += res.get("clicked", [])
        merged["text"] += res.get("text", "")
        merged["isSuccess"] = merged.get("isSuccess") or res.get("isSuccess")
        merged["isError"] = merged.get("isError") or res.get("isError")
        merged["isStockWarning"] = merged.get("isStockWarning") or res.get("isStockWarning")
        if log and res.get("clicked"):
            kind = "WARN" if res.get("isStockWarning") else "INFO"
            log(f"Pop-up terdeteksi -> klik: {', '.join(res['clicked'])}", kind)
        cur = res
    return merged, clicked


# ============================================================
# SIMPAN
# ============================================================
def execute_save_draft(driver):
    return driver.execute_script("""
        function triggerClick(el) {
            if (!el || !el.isConnected) return;
            el.focus?.();
            ['pointerdown','mousedown','pointerup','mouseup','click'].forEach(evt => {
                try { el.dispatchEvent(new MouseEvent(evt, {bubbles:true, cancelable:true, view:window, buttons:1})); } catch(e) {}
            });
            if (typeof el.click === 'function') { try { el.click(); } catch(e) {} }
        }
        function processDoc(doc) {
            let toggles = doc.querySelectorAll('.button-dropdown .dropdown-toggle, .dropdown-toggle');
            toggles.forEach(t => { try { triggerClick(t); } catch(e) {} });
            let elements = Array.from(doc.querySelectorAll('span, a, li'));
            let targetSpan = elements.find(el => (el.innerText || el.textContent || '').trim() === 'Sebagai Draf');
            if (targetSpan) {
                let targetAnchor = targetSpan.closest('a') || targetSpan;
                triggerClick(targetAnchor);
                try {
                    let win = doc.defaultView || window;
                    if (win.ko && win.ko.dataFor) {
                        let data = win.ko.dataFor(targetAnchor);
                        if (data && typeof data.click === 'function') data.click();
                    }
                } catch(e) {}
                return true;
            }
            return false;
        }
        if (processDoc(document)) return true;
        try { if (processDoc(window.top.document)) return true; } catch(e) {}
        let iframes = document.querySelectorAll('iframe');
        for (let f of iframes) { try { if (f.contentDocument && processDoc(f.contentDocument)) return true; } catch(e) {} }
        return false;
    """)


def execute_save_final(driver):
    return driver.execute_script("""
        function triggerClick(el) {
            if (!el || !el.isConnected) return;
            el.focus?.();
            ['pointerdown','mousedown','pointerup','mouseup','click'].forEach(evt => {
                try { el.dispatchEvent(new MouseEvent(evt, {bubbles:true, cancelable:true, view:window, buttons:1})); } catch(e) {}
            });
            if (typeof el.click === 'function') { try { el.click(); } catch(e) {} }
        }
        function processDoc(doc) {
            let toggles = doc.querySelectorAll('.button-dropdown .dropdown-toggle, .dropdown-toggle');
            toggles.forEach(t => { try { triggerClick(t); } catch(e) {} });
            let elements = Array.from(doc.querySelectorAll('span, a, li'));
            let target = elements.find(el => (el.innerText || el.textContent || '').trim() === 'Simpan');
            if (target) {
                let targetAnchor = target.closest('a') || target;
                triggerClick(targetAnchor);
                try {
                    let win = doc.defaultView || window;
                    if (win.ko && win.ko.dataFor) {
                        let data = win.ko.dataFor(targetAnchor);
                        if (data && typeof data.click === 'function') data.click();
                    }
                } catch(e) {}
                return true;
            }
            return false;
        }
        if (processDoc(document)) return true;
        try { if (processDoc(window.top.document)) return true; } catch(e) {}
        let iframes = document.querySelectorAll('iframe');
        for (let f of iframes) { try { if (f.contentDocument && processDoc(f.contentDocument)) return true; } catch(e) {} }
        return false;
    """)


# ============================================================
# PROSES 1 FILE (flow terbukti)
# ============================================================
def process_single_file(driver, file_path, global_idx, total_files, item, date_str, mode, log):
    filename = os.path.basename(file_path)
    log(f"──── File [{global_idx}/{total_files}]: {filename} ────", "FILE")
    log(f"COA {item['account']} | Cabang {item['branch']} | Tanggal {date_str} | Mode {mode}")
    log(f"Keterangan: {item['memo']}")

    log("[1/5] Mengisi tanggal transaksi")
    got = set_transaction_date(driver, date_str)
    if got != date_str:
        human_pause(0.6, 1.2)
        got = set_transaction_date(driver, date_str)
    if got != date_str:
        log(f"Tanggal tidak terverifikasi (terbaca: {got}, target: {date_str}) -> file digagalkan.", "ERROR")
        return False
    human_pause(0.6, 1.2)

    log("[2/5] Mengisi Info Lainnya (Akun, Keterangan, Cabang)")
    verify_res = fill_info_lainnya(driver, item["account"], item["memo"], item["branch"], log)
    acc_ok = bool(verify_res.get("accountFound")) and bool(verify_res.get("accountMatch"))
    br_ok = bool(verify_res.get("branchFound")) and bool(verify_res.get("branchSet"))
    log(f"Akun   : {'OK' if acc_ok else 'GAGAL'} [{verify_res.get('accountNo') or '-'}]", "SUCCESS" if acc_ok else "ERROR")
    log(f"Cabang : {'OK' if br_ok else 'GAGAL'} [{verify_res.get('branchName') or '-'}]", "SUCCESS" if br_ok else "ERROR")
    if not acc_ok or not br_ok:
        log("Verifikasi akun/cabang gagal -> file digagalkan.", "ERROR")
        return False
    human_pause(0.6, 1.2)

    log("[3/5] Mengunggah file Excel")
    try:
        trigger_excel_import(driver, file_path)
    except Exception as e:
        log(f"Gagal upload: {e}", "ERROR")
        return False

    log("[4/5] Menunggu konfirmasi hasil impor")
    popup_eval = wait_import_result(driver, timeout=45, log=log)
    if popup_eval is None:
        log("Pop-up hasil impor tidak muncul sampai batas waktu -> file dilewati.", "ERROR")
        return False
    if popup_eval.get("isError"):
        log(f"DITOLAK Accurate: {popup_eval.get('text','')}", "ERROR")
        human_pause(1.5, 2.5)
        return False
    if popup_eval.get("isSuccess"):
        log("Accurate mengonfirmasi impor berhasil.", "SUCCESS")
    else:
        log(f"Pop-up ditutup: {popup_eval.get('text','')[:70]}")
    human_pause(1.0, 2.0)

    if mode == MODE_IMPORT:
        log("[5/5] Mode Hanya Import: transaksi dibiarkan terbuka (tidak disimpan).", "WARN")
        human_pause(1.5, 2.5)
        return True

    if mode == MODE_DRAFT:
        log("[5/5] Menyimpan sebagai draf")
        ok = False
        for _ in range(3):
            if execute_save_draft(driver):
                ok = True
                break
            human_pause(0.8, 1.4)
        if not ok:
            log("Perintah simpan draf GAGAL terkirim.", "ERROR")
            return False
        log("Perintah simpan draf terkirim.", "SUCCESS")
    else:
        log("[5/5] Menyimpan transaksi (Approve)")
        ok = False
        for _ in range(3):
            if execute_save_final(driver):
                ok = True
                break
            human_pause(0.8, 1.4)
        if not ok:
            log("Perintah simpan final GAGAL terkirim.", "ERROR")
            return False
        log("Perintah simpan final terkirim.", "SUCCESS")

    human_pause(1.0, 2.2)
    log("[+] Menunggu pop-up peringatan stok pasca-simpan (kondisional)")
    popup2, clicked2 = handle_popups_flow(driver, wait_rounds=10, drain_rounds=6, delay=0.9, log=log)
    if any(c == "LANJUTKAN" for c in clicked2):
        log("Peringatan stok pasca-simpan -> LANJUTKAN ditekan otomatis.", "WARN")
    if popup2 and popup2.get("isError"):
        log(f"Permasalahan pasca-simpan: {popup2.get('text','')}", "ERROR")
        return False

    human_pause(2.0, 3.5)
    return True